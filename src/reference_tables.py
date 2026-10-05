"""Read local reference tables and select text without evaluating workbook cells."""

import csv
from io import BytesIO, StringIO
from pathlib import Path
from threading import Lock
import unicodedata
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.evaluation import MAX_REFERENCE_BYTES, parse_reference

MAX_UNCOMPRESSED_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 100
_CSV_LOCK = Lock()

_TRADITIONAL_HEADERS = {
    "textzhtw", "textzhhant", "transcriptzhtw", "transcriptzhhant",
    "zhtw", "zhhant", "traditionalchinese", "chinesetraditional",
    "traditionalchinesetext", "繁體中文", "繁體", "繁體文字", "繁體文本",
    "繁體逐字稿", "正體中文", "正體",
}
_SOURCE_HEADERS = {
    "text", "transcript", "transcription", "utterance", "content", "source",
    "sourcetext", "sourcetranscript", "originaltext", "original", "verbatim",
    "sourceen", "sourceenglish", "原文", "逐字稿", "逐字文本", "文本", "文字",
    "內容", "内容", "轉錄", "转录", "台語", "臺語", "台語文本", "中文",
    "chinese", "simplifiedchinese", "chinesesimplified", "簡體中文", "简体中文",
    "zh", "zhcn", "zhhans",
} | {
    prefix + language
    for prefix in ("text", "transcript", "utterance")
    for language in ("zh", "zhcn", "zhhans", "chinese", "ja", "japanese", "ko", "korean",
                     "fr", "french", "es", "spanish", "de", "german", "ar", "arabic")
}
_ENGLISH_HEADERS = {
    "translationen", "translationenglish", "englishtranslation", "translation",
    "referenceen", "enreference", "englishreference", "referenceenglish",
    "english", "en", "texten", "textenglish", "translatedtext", "targettext",
    "target", "英文", "英語", "英语", "英文翻譯", "英文翻译", "英語翻譯", "英语翻译",
}
_STRUCTURAL_HEADERS = {
    "id", "start", "starts", "starttime", "end", "ends", "endtime", "speaker",
    "speakerid", "timestamp", "time", "overlap", "backchannel", "notes", "note",
}


class _TableInputError(ValueError):
    """Keep actionable validation errors separate from library parsing failures."""


def _populated(value: str | None) -> bool:
    return value is None or bool(value.strip())


def _trim_grid(rows: list[list[str | None]]) -> list[list[str | None]]:
    while rows and not any(_populated(value) for value in rows[-1]):
        rows.pop()
    width = max((index + 1 for row in rows for index, value in enumerate(row) if _populated(value)), default=0)
    return [row[:width] + [""] * max(0, width - len(row)) for row in rows]


def _check_shape(row_number: int, width: int) -> None:
    if row_number > MAX_ROWS:
        raise _TableInputError("Each reference table must have at most 10,000 rows, including headers.")
    if width > MAX_COLUMNS:
        raise _TableInputError("Each reference table must have at most 100 columns.")


def _read_xlsx(data: bytes) -> dict[str, list[list[str | None]]]:
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise _TableInputError("The XLSX contents must be 10 MiB or smaller when uncompressed.")
            if any(entry.flag_bits & 1 for entry in entries):
                raise _TableInputError("Password-protected reference workbooks are not supported.")
            content_types = archive.read("[Content_Types].xml").lower()
            if b"macroenabled" in content_types or any("vbaproject" in entry.filename.lower() for entry in entries):
                raise _TableInputError("Macro-enabled workbooks are not supported. Save a plain XLSX copy.")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            tables = {}
            for sheet in workbook.worksheets:
                _check_shape(sheet.max_row or 0, sheet.max_column or 0)
                # Some exporters write incorrect dimensions; also check actual cells.
                sheet.reset_dimensions()
                rows = []
                for row_number, cells in enumerate(sheet.iter_rows(), start=1):
                    _check_shape(row_number, len(cells))
                    rows.append([
                        None if cell.data_type in {"f", "e"}
                        else "" if cell.value is None else str(cell.value)
                        for cell in cells
                    ])
                rows = _trim_grid(rows)
                if rows:
                    tables[sheet.title] = rows
            return tables
        finally:
            workbook.close()
    except _TableInputError:
        raise
    except Exception as exc:
        raise ValueError("Could not read this XLSX file. It may be invalid or corrupted.") from exc


def _read_delimited(data: bytes, *, tsv: bool) -> list[list[str | None]]:
    try:
        encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        text = data.decode(encoding)
    except UnicodeDecodeError as exc:
        raise ValueError("CSV and TSV references must use UTF-8 or BOM-marked UTF-16 text.") from exc
    if any(unicodedata.category(character) == "Cc" and character not in "\t\r\n" for character in text):
        raise ValueError("The reference table contains binary or unsupported control characters.")
    delimiter = "\t" if tsv else ","
    if not tsv:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
        except csv.Error:
            pass
    try:
        with _CSV_LOCK:
            previous_limit = csv.field_size_limit()
            try:
                csv.field_size_limit(MAX_REFERENCE_BYTES)
                rows = []
                for row_number, row in enumerate(csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True), start=1):
                    _check_shape(row_number, len(row))
                    rows.append(row)
            finally:
                csv.field_size_limit(previous_limit)
        return _trim_grid(rows)
    except csv.Error as exc:
        raise ValueError("Could not read this CSV or TSV file. Check its separators and quoted fields.") from exc


def read_reference_tables(data: bytes, filename: str) -> dict[str, list[list[str | None]]]:
    """Read XLSX, CSV, or TSV grids in memory; never evaluate formulas or links.

    Strings preserve cell values; blanks are ''. XLSX formulas and error cells are
    None, so unrelated formula columns remain usable but cannot become reference
    text. Preserve leading/internal blank rows and columns; trim only trailing
    emptiness. Ignore empty sheets. Limits: 1 MiB uploaded, 10 MiB uncompressed XLSX,
    and 10,000 rows / 100 columns per sheet. Unsupported or malformed input raises
    ValueError; limits never silently truncate data. CSV formula-like text is literal.
    """
    if len(data) > MAX_REFERENCE_BYTES:
        raise ValueError("The reference table must be 1 MiB or smaller.")
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        tables = _read_xlsx(data)
    elif suffix in {".csv", ".tsv"}:
        rows = _read_delimited(data, tsv=suffix == ".tsv")
        tables = {"TSV" if suffix == ".tsv" else "CSV": rows} if rows else {}
    else:
        raise ValueError("Upload a plain XLSX, CSV, or TSV reference table.")
    if not tables:
        raise ValueError("The reference table contains no usable cells.")
    return tables


def _header_key(value: str | None) -> str:
    return "".join(character for character in unicodedata.normalize("NFKC", value or "").casefold() if character.isalnum())


def suggest_table(tables: dict[str, list[list[str | None]]]) -> str:
    """Prefer transcript/transcription sheets, then sheets unrelated to metadata."""
    if not tables:
        raise ValueError("There are no reference tables to select.")
    for name in tables:
        if _header_key(name) in {"transcript", "transcription"}:
            return name
    for name in tables:
        if not any(word in _header_key(name) for word in ("background", "speaker", "note", "metadata")):
            return name
    return next(iter(tables))


def suggest_header_row(rows: list[list[str | None]]) -> int:
    """Return the earliest best header match in the first 20 rows, or zero."""
    known = _TRADITIONAL_HEADERS | _SOURCE_HEADERS | _ENGLISH_HEADERS | _STRUCTURAL_HEADERS
    scores = [sum(_header_key(value) in known for value in row) for row in rows[:20]]
    return max(range(len(scores)), key=scores.__getitem__) if scores else 0


def _header(rows: list[list[str | None]], header_row: int | None) -> list[str | None]:
    if header_row is None:
        return []
    if not 0 <= header_row < len(rows):
        raise ValueError("Choose a header row that exists in the selected table.")
    return rows[header_row]


def table_columns(rows: list[list[str | None]], header_row: int | None) -> list[str]:
    """Return unique column-letter labels, optionally including header text."""
    header = _header(rows, header_row)
    labels = []
    for index in range(max(map(len, rows), default=0)):
        value = header[index] if index < len(header) else ""
        name = "(formula or error)" if value is None else " ".join(value.split())
        labels.append(get_column_letter(index + 1) + (f": {name}" if name else ""))
    return labels


def recommend_columns(rows: list[list[str | None]], header_row: int | None) -> dict[str, int | None]:
    """Suggest source/English columns by headers, leaving ambiguous choices unset.

    Traditional Chinese takes source precedence over other source aliases. An
    English translation column is never inferred as source. ``display`` is the
    English reference column; callers may fall back to source when it is absent.
    No content is guessed from rows when headers are disabled.
    """
    keys = [_header_key(value) for value in _header(rows, header_row)]
    source = [index for index, key in enumerate(keys) if key in _TRADITIONAL_HEADERS]
    if not source:
        source = [index for index, key in enumerate(keys) if key in _SOURCE_HEADERS]
    display = [index for index, key in enumerate(keys) if key in _ENGLISH_HEADERS]
    return {
        "source": source[0] if len(source) == 1 else None,
        "display": display[0] if len(display) == 1 else None,
    }


def parsed_table_reference(rows: list[list[str | None]], *, header_row: int | None,
                           text_column: int, format: str = "auto") -> dict:
    """Parse only a selected column, retaining its original row numbers and text.

    Skip rows through the selected header and blank selected cells. Formula/error
    cells in the selected data column raise ValueError asking for pasted values.
    Other columns never enter the reference. The parse_reference result additionally
    contains ``rows``: [{'row_number': original 1-based row, 'text': raw cell text}].
    This preserves row order; it does not claim audio-to-row timing alignment.
    """
    _header(rows, header_row)
    if not 0 <= text_column < max(map(len, rows), default=0):
        raise ValueError("Choose a text column that exists in the selected table.")
    selected = []
    for index in range(0 if header_row is None else header_row + 1, len(rows)):
        value = rows[index][text_column] if text_column < len(rows[index]) else ""
        if value is None:
            cell = f"{get_column_letter(text_column + 1)}{index + 1}"
            raise ValueError(f"Cell {cell} contains a formula or error. Paste its text as a value and upload again.")
        if value.strip():
            selected.append({"row_number": index + 1, "text": value})
    if not selected:
        raise ValueError("The selected text column contains no reference text below the header.")
    result = parse_reference("\n".join(row["text"] for row in selected).encode("utf-8"), format=format)
    result["rows"] = selected
    return result
