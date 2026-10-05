"""Exercise table imports using synthetic in-memory workbooks and delimited text."""

from io import BytesIO
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from src.reference_tables import (
    MAX_COLUMNS, MAX_REFERENCE_BYTES, MAX_ROWS, MAX_UNCOMPRESSED_BYTES,
    parsed_table_reference, read_reference_tables, recommend_columns,
    suggest_header_row, suggest_table, table_columns,
)


def workbook_bytes(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    data = BytesIO()
    workbook.save(data)
    workbook.close()
    return data.getvalue()


class ReadReferenceTablesTests(unittest.TestCase):
    def test_workbook_preserves_sheet_order_blank_row_positions_and_values(self):
        data = workbook_bytes({
            "transcript": [[], ["Title"], [], ["id", "text", "flag"], [1, "你好", True], [], [2, "Hello", False], []],
            "background_speech": [["text"], ["Unrelated background"]],
            "empty": [],
        })

        tables = read_reference_tables(data, "reference.XLSX")

        self.assertEqual(list(tables), ["transcript", "background_speech"])
        self.assertEqual(tables["transcript"], [
            ["", "", ""], ["Title", "", ""], ["", "", ""],
            ["id", "text", "flag"], ["1", "你好", "True"],
            ["", "", ""], ["2", "Hello", "False"],
        ])

    def test_formula_and_error_cells_are_markers_not_evaluated_text(self):
        data = workbook_bytes({"transcript": [["text", "formula", "error"], ["Hello", "=1+1", "#DIV/0!"]]})

        rows = read_reference_tables(data, "reference.xlsx")["transcript"]

        self.assertEqual(rows, [["text", "formula", "error"], ["Hello", None, None]])
        self.assertEqual(parsed_table_reference(rows, header_row=0, text_column=0)["text"], "Hello")
        for column in (1, 2):
            with self.subTest(column=column):
                with self.assertRaisesRegex(ValueError, "contains a formula or error.*Paste"):
                    parsed_table_reference(rows, header_row=0, text_column=column)

    def test_formula_only_sheet_remains_selectable_for_actionable_validation(self):
        rows = read_reference_tables(workbook_bytes({"sheet": [["=1+1"]]}), "reference.xlsx")["sheet"]

        self.assertEqual(rows, [[None]])
        with self.assertRaisesRegex(ValueError, "Cell A1"):
            parsed_table_reference(rows, header_row=None, text_column=0)

    def test_csv_preserves_quotes_newlines_leading_and_internal_empty_rows(self):
        tables = read_reference_tables(b'\nid,text\n1,"Hello,\nworld"\n\n2,"She said ""yes""."\n\n', "reference.csv")

        self.assertEqual(tables, {"CSV": [
            ["", ""], ["id", "text"], ["1", "Hello,\nworld"],
            ["", ""], ["2", 'She said "yes".'],
        ]})

    def test_csv_semicolon_and_tsv_separators(self):
        for filename, data, name in (
            ("reference.csv", b"text;translation_en\nhello;hello", "CSV"),
            ("reference.tsv", b"text\ttranslation_en\nhello\thello", "TSV"),
        ):
            with self.subTest(filename=filename):
                self.assertEqual(read_reference_tables(data, filename), {
                    name: [["text", "translation_en"], ["hello", "hello"]],
                })

    def test_utf8_and_utf16_bom_are_supported_without_losing_blank_rows(self):
        for encoding in ("utf-8-sig", "utf-16"):
            with self.subTest(encoding=encoding):
                rows = read_reference_tables("\ntext\n你好".encode(encoding), "reference.csv")["CSV"]
                self.assertEqual(rows, [[""], ["text"], ["你好"]])

    def test_large_csv_cell_within_upload_limit_is_supported(self):
        text = "x" * 150_000

        rows = read_reference_tables(("text\n" + text).encode(), "reference.csv")["CSV"]

        self.assertEqual(rows[1][0], text)

    def test_csv_formula_like_strings_remain_literal(self):
        rows = read_reference_tables(b'text\n"=SUM(A1:A2)"', "reference.csv")["CSV"]

        self.assertEqual(rows[1][0], "=SUM(A1:A2)")

    def test_invalid_encoding_binary_and_quoting_have_clear_errors(self):
        for data, message in (
            (b"text\ncaf\xe9", "UTF-8"),
            (b"text\nhello\x00world", "binary"),
            (b'text\n"unclosed', "quoted fields"),
        ):
            with self.subTest(data=data):
                with self.assertRaisesRegex(ValueError, message):
                    read_reference_tables(data, "reference.csv")

    def test_empty_tables_are_rejected(self):
        for data, filename in (
            (b"\n,\n,\n", "reference.csv"),
            (workbook_bytes({"empty": []}), "reference.xlsx"),
        ):
            with self.subTest(filename=filename):
                with self.assertRaisesRegex(ValueError, "no usable cells"):
                    read_reference_tables(data, filename)

    def test_corrupt_and_unsupported_workbooks_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "invalid or corrupted"):
            read_reference_tables(b"not a workbook", "reference.xlsx")
        for filename in ("reference.xls", "reference.xlsm", "reference.txt"):
            with self.subTest(filename=filename):
                with self.assertRaisesRegex(ValueError, "plain XLSX, CSV, or TSV"):
                    read_reference_tables(b"hello", filename)

    def test_renamed_macro_workbook_is_rejected(self):
        original = workbook_bytes({"transcript": [["text"], ["Hello"]]})
        modified = BytesIO()
        with ZipFile(BytesIO(original)) as source, ZipFile(modified, "w", ZIP_DEFLATED) as target:
            for entry in source.infolist():
                target.writestr(entry.filename, source.read(entry.filename))
            target.writestr("xl/vbaProject.bin", b"synthetic nonexecutable test marker")

        with self.assertRaisesRegex(ValueError, "Macro-enabled"):
            read_reference_tables(modified.getvalue(), "renamed.xlsx")

    def test_upload_and_uncompressed_limits(self):
        with self.assertRaisesRegex(ValueError, "1 MiB"):
            read_reference_tables(b"x" * (MAX_REFERENCE_BYTES + 1), "reference.csv")
        expanded = BytesIO()
        with ZipFile(expanded, "w", ZIP_DEFLATED) as archive:
            archive.writestr("oversized.xml", b"x" * (MAX_UNCOMPRESSED_BYTES + 1))
        with self.assertRaisesRegex(ValueError, "10 MiB"):
            read_reference_tables(expanded.getvalue(), "reference.xlsx")

    def test_csv_grid_limits_do_not_silently_truncate(self):
        for data, message in (
            (("x\n" * (MAX_ROWS + 1)).encode(), "10,000 rows"),
            ((",".join(["x"] * (MAX_COLUMNS + 1))).encode(), "100 columns"),
        ):
            with self.subTest(message=message):
                with self.assertRaisesRegex(ValueError, message):
                    read_reference_tables(data, "reference.csv")

    def test_xlsx_grid_limits_include_sparse_faraway_cells(self):
        for cell, message in ((f"A{MAX_ROWS + 1}", "10,000 rows"), ("CW1", "100 columns")):
            with self.subTest(cell=cell):
                workbook = Workbook()
                workbook.active[cell] = "out of bounds"
                data = BytesIO()
                workbook.save(data)
                workbook.close()
                with self.assertRaisesRegex(ValueError, message):
                    read_reference_tables(data.getvalue(), "reference.xlsx")


class TableSelectionTests(unittest.TestCase):
    def test_suggest_table_prefers_transcript_and_avoids_background_metadata(self):
        self.assertEqual(suggest_table({"background_speech": [], "speakers": [], "Transcript": [], "Sheet1": []}), "Transcript")
        self.assertEqual(suggest_table({"notes": [], "background_speech": [], "Sheet1": []}), "Sheet1")
        self.assertEqual(suggest_table({"speakers": []}), "speakers")
        with self.assertRaisesRegex(ValueError, "no reference tables"):
            suggest_table({})

    def test_header_suggestion_skips_title_and_blank_rows_and_is_bounded(self):
        rows = [["Meeting export"], [], ["id", "speaker", "text_zh_TW", "translation_en"], ["1", "Alice", "你好", "Hello"]]

        self.assertEqual(suggest_header_row(rows), 2)
        self.assertEqual(suggest_header_row([["unrecognized"]] * 20 + [rows[2]]), 0)
        self.assertEqual(suggest_header_row([]), 0)

    def test_column_labels_are_unique_for_duplicate_or_missing_headers(self):
        rows = [["text", "text", "", None], ["a", "b", "c", "d"]]

        self.assertEqual(table_columns(rows, 0), ["A: text", "B: text", "C", "D: (formula or error)"])
        self.assertEqual(table_columns(rows, None), ["A", "B", "C", "D"])

    def test_source_prefers_traditional_chinese_and_display_uses_english_reference(self):
        headers = ["id", "start_s", "end_s", "speaker_id", "speaker", "overlap", "backchannel", "text_zh_TW", "text_zh_CN", "translation_en", "notes"]

        self.assertEqual(recommend_columns([headers], 0), {"source": 7, "display": 9})
        self.assertEqual(recommend_columns([["繁體中文", "英語翻譯"]], 0), {"source": 0, "display": 1})

    def test_ambiguous_missing_or_disabled_headers_require_a_choice(self):
        for rows, header_row, expected in (
            ([["text", "transcript", "notes"]], 0, {"source": None, "display": None}),
            ([["text_zh_TW", "繁體中文", "english", "translation_en"]], 0, {"source": None, "display": None}),
            ([["translation_en"]], 0, {"source": None, "display": 0}),
            ([["你好", "Hello"]], 0, {"source": None, "display": None}),
            ([["text", "translation_en"]], None, {"source": None, "display": None}),
        ):
            with self.subTest(rows=rows, header_row=header_row):
                self.assertEqual(recommend_columns(rows, header_row), expected)

    def test_selected_column_excludes_other_columns_preamble_and_header(self):
        rows = [
            ["Meeting title"], [], ["id", "source", "translation_en", "notes"],
            ["1", " 你好 ", "Hello", "Do not score this"],
            ["2", "", "English only", "metadata"],
            ["3", "謝謝", "Thank you", None],
        ]

        source = parsed_table_reference(rows, header_row=2, text_column=1)
        english = parsed_table_reference(rows, header_row=2, text_column=2)

        self.assertEqual(source["text"], "你好 \n謝謝")
        self.assertEqual(source["rows"], [{"row_number": 4, "text": " 你好 "}, {"row_number": 6, "text": "謝謝"}])
        self.assertEqual(english["text"], "Hello\nEnglish only\nThank you")
        self.assertTrue(source["has_cjk"])
        self.assertFalse(english["has_cjk"])

    def test_headerless_and_ragged_tables_keep_original_row_numbers(self):
        result = parsed_table_reference([["Hello"], [], ["World", "other"]], header_row=None, text_column=0)

        self.assertEqual(result["text"], "Hello\nWorld")
        self.assertEqual(result["rows"], [{"row_number": 1, "text": "Hello"}, {"row_number": 3, "text": "World"}])

    def test_selected_cells_can_use_existing_transcript_cleanup(self):
        rows = [["text"], ["[00:01] SPK1: (overlap) 你好"], ["[00:02] SPK2: Hello"]]

        result = parsed_table_reference(rows, header_row=0, text_column=0)

        self.assertEqual(result["text"], "你好\nHello")
        self.assertEqual(result["segment_count"], 2)
        self.assertEqual(result["rows"][0]["text"], "[00:01] SPK1: (overlap) 你好")

    def test_invalid_header_column_and_empty_selection_fail_actionably(self):
        rows = [["text", "other"], ["", "content"]]
        with self.assertRaisesRegex(ValueError, "header row"):
            parsed_table_reference(rows, header_row=3, text_column=0)
        with self.assertRaisesRegex(ValueError, "text column that exists"):
            parsed_table_reference(rows, header_row=0, text_column=5)
        with self.assertRaisesRegex(ValueError, "no reference text"):
            parsed_table_reference(rows, header_row=0, text_column=0)


if __name__ == "__main__":
    unittest.main()
